"""
export_excel.py
===============
Xuất kết quả nội lực ra file Excel (.xlsx) với định dạng chuyên nghiệp.

Tương đương VBA: Phần ghi sheet + tô màu trong Module1, Module2, Module4.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional, Union

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from core_engine import (
    ReactionRow,
    ResultRow,
    ResultRowMulti,
    SummaryRow,
    get_unique_locations,
)


SHEET_NOILUC = "NoiLuc_NhieuViTri"
SHEET_TORSION = "torsional purlin check"
NAME_VITRI_LIST = "TorsionalPurlinViTriList"

FILL_TORSION_H = PatternFill("solid", fgColor="C6EFCE")  # RGB(198, 239, 206)
FILL_TORSION_M3 = PatternFill("solid", fgColor="FFC7CE")  # RGB(255, 199, 206)


# ── Màu sắc (giống VBA gốc) ─────────────────────────────────

FILL_P = PatternFill("solid", fgColor="BDD7EE")       # Xanh dương nhạt
FILL_V2 = PatternFill("solid", fgColor="92D050")      # Xanh lá nhạt
FILL_V3 = PatternFill("solid", fgColor="FCE4D6")      # Cam/hồng nhạt
FILL_M2 = PatternFill("solid", fgColor="FFF2CC")      # Vàng nhạt
FILL_M3 = PatternFill("solid", fgColor="F8CBAD")      # Cam đậm
FILL_HEADER = PatternFill("solid", fgColor="4472C4")   # Xanh đậm header

# Excel InternalForce: tô ô cực trị theo dấu
FILL_POS = PatternFill("solid", fgColor="E2EFDA")      # RGB(226, 239, 218)
FILL_NEG = PatternFill("solid", fgColor="FCE4D6")      # RGB(252, 228, 214)
FONT_POS = Font(bold=True, color="385723")             # RGB(56, 87, 35)
FONT_NEG = Font(bold=True, color="C00000")             # RGB(192, 0, 0)

FONT_HEADER = Font(bold=True, color="FFFFFF", size=11)
FONT_BOLD = Font(bold=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
CENTER = Alignment(horizontal="center", vertical="center")


def _auto_fit_columns(ws):
    """Tự điều chỉnh độ rộng cột theo nội dung."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                cell_len = len(str(cell.value))
                if cell_len > max_len:
                    max_len = cell_len
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)


def _style_header(ws, row_num: int, col_count: int):
    """Định dạng dòng header."""
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row_num, column=c)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def _mark_extreme_cell(cell):
    """Excel InternalForce: ô cực trị in đậm, >= 0 xanh lá / < 0 cam chữ đỏ."""
    try:
        positive = float(cell.value) >= 0
    except (TypeError, ValueError):
        positive = True
    cell.fill = FILL_POS if positive else FILL_NEG
    cell.font = FONT_POS if positive else FONT_NEG


def _dim(value: float):
    """Excel: kích thước = 0 → 'N/A'."""
    return value if value > 0 else "N/A"


# ═══════════════════════════════════════════════════════════════
# XUẤT KẾT QUẢ MODULE 1: NỘI LỰC NÚT ĐẦU/CUỐI (Excel InternalForce)
# ═══════════════════════════════════════════════════════════════

NODE_HEADERS = [
    "Tên Thanh", "Tiết Diện", "Chiều Dài (m)", "Vị Trí Nút",
    "Loại Cực Trị", "Nguồn Tải (Combo/Case)", "Station (m)",
    "P (kN)", "M3 (kNm)", "M2 (kNm)", "T (kNm)", "V2 (kN)", "V3 (kN)",
    "Loại TD", "t3 (m)", "t2 (m)", "tf (m)", "tw (m)", "t2b (m)", "tfb (m)",
    "dis (m)", "r (m)", "Mirror",
]

# Loại cực trị → cột nội lực được tô (số cột trong NODE_HEADERS, tính từ 1)
NODE_EXTREME_COL = {
    "Max P": 8, "Min P": 8,
    "Max M3": 9, "Min M3": 9,
    "Max M2": 10, "Min M2": 10,
    "Max T (Abs)": 11,
    "Max V2 (Abs)": 12,
    "Max V3 (Abs)": 13,
}


def export_node_results(
    data: list[ResultRow],
    filepath: Union[str, Path],
    sheet_name: str = "NoiLuc_Bao_TuongUng",
):
    """Xuất bảng nội lực cực trị nút đầu/cuối ra Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    for c, h in enumerate(NODE_HEADERS, 1):
        ws.cell(row=1, column=c, value=h)
    _style_header(ws, 1, len(NODE_HEADERS))

    for i, row in enumerate(data, 2):
        vals = [
            row.frame_name, row.section_name, row.length, row.node_label,
            row.extreme_type, row.load_case, round(row.obj_station, 3),
            row.P, row.M3, row.M2, row.T, row.V2, row.V3,
            row.sec_type,
            _dim(row.t3), _dim(row.t2), _dim(row.tf), _dim(row.tw),
            _dim(row.t2b), _dim(row.tfb), _dim(row.dis), _dim(row.fillet),
            row.mirror,
        ]
        for c, v in enumerate(vals, 1):
            ws.cell(row=i, column=c, value=v).border = THIN_BORDER

        col = NODE_EXTREME_COL.get(row.extreme_type)
        if col:
            _mark_extreme_cell(ws.cell(row=i, column=col))

    _auto_fit_columns(ws)
    ws.freeze_panes = "A2"
    wb.save(str(filepath))
    return str(filepath)


# ═══════════════════════════════════════════════════════════════
# XUẤT KẾT QUẢ MODULE 2: TÓM TẮT CỰC TRỊ TOÀN CỤC
# ═══════════════════════════════════════════════════════════════

# ═══════════════════════════════════════════════════════════════
# XUẤT PHẢN LỰC NÚT (Excel Reaction)
# ═══════════════════════════════════════════════════════════════

REACTION_HEADERS = [
    "Nút", "Cực Trị", "Nguồn Tải (Combo/Case)", "Step",
    "F1 (kN)", "F2 (kN)", "F3 (kN)", "M1 (kNm)", "M2 (kNm)", "M3 (kNm)",
]
_REACTION_COL = {"F1": 5, "F2": 6, "F3": 7, "M1": 8, "M2": 9, "M3": 10}


def export_reaction_results(
    data: list[ReactionRow],
    filepath: Union[str, Path],
    sheet_name: str = "PhanLuc",
):
    """Xuất bảng phản lực (đúng kiểu lọc đang xem) — tô ô quyết định theo dấu."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    for c, h in enumerate(REACTION_HEADERS, 1):
        ws.cell(row=1, column=c, value=h)
    _style_header(ws, 1, len(REACTION_HEADERS))

    for i, row in enumerate(data, 2):
        vals = [
            row.joint, row.criteria, row.load_case, row.step_text,
            row.F1, row.F2, row.F3, row.M1, row.M2, row.M3,
        ]
        for c, v in enumerate(vals, 1):
            ws.cell(row=i, column=c, value=v).border = THIN_BORDER
        col = _REACTION_COL.get(row.comp)
        if col:
            _mark_extreme_cell(ws.cell(row=i, column=col))

    _auto_fit_columns(ws)
    ws.freeze_panes = "A2"
    wb.save(str(filepath))
    return str(filepath)


SUMMARY_EXTREME_COL = {
    "Mmax": "M33 (kNm)", "Mmin": "M33 (kNm)",
    "V2max": "V2 (kN)",
    "Nmax": "Axial (kN)", "Nmin": "Axial (kN)",
    "V3max": "V3 (kN)",
}


def export_summary(
    summary: list[SummaryRow],
    filepath: Union[str, Path],
    columns_order: Optional[list[str]] = None,
    condition_order: Optional[list[str]] = None,
    scope_text: str = "",
):
    """
    Xuất bảng tóm tắt cực trị ra Excel — 1 sheet duy nhất (đã loại bỏ cột Thanh).

    columns_order / condition_order: đúng thứ tự cột và các hàng đang hiển thị
    trên tool (dòng đang ẩn thì không xuất) để dán RAM Connection không bị lệch.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "TomTat_CucTri"

    rows = summary
    if condition_order:
        by_cond = {r.condition: r for r in summary}
        rows = [by_cond[c] for c in condition_order if c in by_cond]

    title = f"--- {scope_text.upper()} ---" if scope_text else "--- TÓM TẮT CỰC TRỊ ---"
    _write_summary_sheet(ws, rows, title, columns_order)

    wb.save(str(filepath))
    return str(filepath)


def _write_summary_sheet(ws, data: list[SummaryRow], title: str, columns_order: Optional[list[str]] = None):
    """Ghi 1 sheet tóm tắt (không có cột Thanh)."""
    ws.cell(row=1, column=1, value=title).font = FONT_BOLD

    default_headers = ["STT", "Điều kiện", "V2 (kN)", "Axial (kN)", "M33 (kNm)", "M22 (kNm)", "V3 (kN)"]
    headers = [c for c in columns_order if c != "Thanh"] if columns_order else default_headers

    for c, h in enumerate(headers, 1):
        ws.cell(row=2, column=c, value=h)
    _style_header(ws, 2, len(headers))

    # Map tên cột tới giá trị (STT đánh lại theo thứ tự hiển thị)
    for i, row in enumerate(data, 3):
        val_map = {
            "STT": i - 2,
            "Điều kiện": row.condition,
            "Điều Kiện": row.condition,
            "V2 (kN)": row.V2,
            "V3 (kN)": row.V3,
            "Axial (kN)": row.axial,
            "M33 (kNm)": row.M33,
            "M22 (kNm)": row.M22,
            "Lấy từ": row.source,
            "Tổ hợp (Combo/Case)": row.load_case,
        }
        for c, h in enumerate(headers, 1):
            val = val_map.get(h, "")
            ws.cell(row=i, column=c, value=val)
            ws.cell(row=i, column=c).border = THIN_BORDER

        target = SUMMARY_EXTREME_COL.get(row.condition)
        if row.source and target in headers:
            _mark_extreme_cell(ws.cell(row=i, column=headers.index(target) + 1))

    _auto_fit_columns(ws)


# ═══════════════════════════════════════════════════════════════
# XUẤT KẾT QUẢ MODULE 4: NHIỀU VỊ TRÍ
# ═══════════════════════════════════════════════════════════════

def export_multi_station_results(
    data: list[ResultRowMulti],
    filepath: Union[str, Path],
    sheet_name: str = SHEET_NOILUC,
):
    """
    Xuất bảng nội lực nhiều vị trí ra Excel.

    Giống VBA Module4: sheet NoiLuc_NhieuViTri + sheet torsional purlin check
    (dropdown ViTri, công thức AGGREGATE tìm M3 MIN).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    headers = [
        "Tên Thanh", "Tiết Diện", "htb (mm)", "tw (mm)", "bf (mm)", "tf (mm)",
        "Chiều Dài (m)", "Vị Trí", "Loại Cực Trị", "Nguồn Tải (Combo/Case)",
        "P (kN)", "V2 (kN)", "V3 (kN)", "M2 (kNm)", "M3 (kNm)",
    ]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    _style_header(ws, 1, len(headers))

    for i, row in enumerate(data, 2):
        vals = [
            row.frame_name, row.section_name, row.htb, row.tw, row.bf, row.tf,
            row.length, row.location_label, row.extreme_type, row.load_case,
            row.P, row.V2, row.V3, row.M2, row.M3,
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.border = THIN_BORDER

    _auto_fit_columns(ws)

    if data:
        _write_torsional_purlin_check(wb, ws, len(data) + 1, data)

    wb.save(str(filepath))
    return str(filepath)


def _xl_sheet(name: str) -> str:
    """Tên sheet Excel có dấu nháy, escape '."""
    return "'" + name.replace("'", "''") + "'"


def _write_torsional_purlin_check(
    wb: Workbook,
    ws_source,
    last_source_row: int,
    data: list[ResultRowMulti],
):
    """VBA: TaoSheetTorsionalPurlinCheck."""
    if SHEET_TORSION in wb.sheetnames:
        del wb[SHEET_TORSION]
    ws = wb.create_sheet(SHEET_TORSION)

    # Xóa named range cũ
    try:
        existing = wb.defined_names.get(NAME_VITRI_LIST)
        if existing is not None:
            del wb.defined_names[NAME_VITRI_LIST]
    except Exception:
        pass

    # A1:O1 tiêu đề
    ws.merge_cells("A1:O1")
    title = ws["A1"]
    title.value = SHEET_TORSION
    title.font = Font(bold=True, size=14)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    ws.merge_cells("A2:O2")
    ws["A2"].value = (
        "Chon ViTri tai o H4 -> tu dong tim M3 MIN va lay dong tuong ung"
    )
    ws["A2"].font = Font(italic=True)

    headers = [
        "Ten Thanh", "Tiet Dien", "htb (mm)", "tw (mm)", "bf (mm)", "tf (mm)",
        "Chieu Dai (m)", "Vi Tri", "Loai Cuc Tri", "Nguon Tai (Combo/Case)",
        "P (kN)", "V2 (kN)", "V3 (kN)", "M2 (kNm)", "M3 (kNm)",
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = FONT_BOLD
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    locations = get_unique_locations(data)
    if not locations:
        return

    for i, loc in enumerate(locations):
        ws.cell(row=4 + i, column=17, value=loc)  # cột Q

    q_last = 3 + len(locations)
    src = _xl_sheet(ws_source.title)
    chk = _xl_sheet(SHEET_TORSION)
    n = last_source_row

    ref = f"{chk}!$Q$4:$Q${q_last}"
    dn = DefinedName(name=NAME_VITRI_LIST, attr_text=ref)
    try:
        wb.defined_names.add(dn)
    except Exception:
        try:
            wb.defined_names.append(dn)
        except Exception:
            wb.defined_names[NAME_VITRI_LIST] = dn

    dv = DataValidation(
        type="list",
        formula1=f"={NAME_VITRI_LIST}",
        allow_blank=False,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="ViTri khong hop le",
        error="Vui long chon ViTri trong danh sach.",
    )
    dv.add(ws["H4"])
    ws.add_data_validation(dv)
    ws["H4"].value = locations[0]

    ws["O4"] = (
        f'=IF($H$4="","",IFERROR(AGGREGATE(15,6,'
        f'{src}!$O$2:$O${n}/({src}!$H$2:$H${n}=$H$4),1),""))'
    )
    ws["P4"] = (
        f'=IF($H$4="","",IFERROR(AGGREGATE(15,6,'
        f'(ROW({src}!$A$2:$A${n})-ROW({src}!$A$2)+1)/'
        f'(({src}!$H$2:$H${n}=$H$4)*({src}!$O$2:$O${n}=$O$4)),1),""))'
    )

    for c in list(range(1, 8)) + list(range(9, 15)):
        col = get_column_letter(c)
        ws.cell(row=4, column=c).value = (
            f'=IF($P$4="","",INDEX({src}!{col}2:{col}{n},$P$4))'
        )

    for c in range(1, 16):
        cell = ws.cell(row=4, column=c)
        cell.font = FONT_BOLD
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="center")

    ws["H4"].fill = FILL_TORSION_H
    ws["H4"].font = FONT_BOLD
    ws["H4"].alignment = CENTER

    ws["O4"].fill = FILL_TORSION_M3
    ws["O4"].font = FONT_BOLD
    ws["O4"].alignment = CENTER
    ws["O4"].number_format = "0.00"

    ws.column_dimensions["P"].hidden = True
    ws.column_dimensions["Q"].hidden = True

    _auto_fit_columns(ws)
    for col in ("C", "D", "E", "F", "H", "I", "J", "K", "L", "M", "N", "O"):
        for cell in ws[f"{col}4:{col}4"]:
            pass
    for c in range(3, 7):
        ws.cell(row=4, column=c).alignment = CENTER
    for c in range(8, 16):
        ws.cell(row=4, column=c).alignment = CENTER
